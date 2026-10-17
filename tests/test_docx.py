"""Word export and import: the round trip, what Word itself writes, what is
refused with which code, and who may do either."""

from __future__ import annotations

import io
import json
import zipfile
from typing import Any
from urllib.parse import unquote

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from httpx import AsyncClient

from api.core.config import get_settings
from api.core.downloads import content_disposition
from api.services.docx_export import export_filename, to_docx
from api.services.docx_format import DOCX_CONTENT_TYPE, attribute_names
from api.services.docx_import import (
    CorruptDocumentError,
    Dropped,
    choose_title,
    import_docx,
)
from tests.conftest import TestUser
from tests.docx_helpers import (
    FIXTURES,
    body_xml,
    compound_file,
    docx_bytes,
    with_extra_part,
    with_part,
)
from tests.editor_fixtures import load, manifest

LIMIT = 100 * 1024 * 1024


def docs_url(workspace: dict[str, Any], document_id: str | None = None) -> str:
    base = f"/workspaces/{workspace['id']}/documents"
    return f"{base}/{document_id}" if document_id else base


def kinds(content: dict[str, Any]) -> list[str]:
    return [block["type"] for block in content["content"]]


def read(data: bytes) -> Any:
    return import_docx(data, max_uncompressed_bytes=LIMIT)


async def upload(
    client: AsyncClient,
    workspace: dict[str, Any],
    user: TestUser,
    data: bytes,
    filename: str = "Report.docx",
) -> Any:
    return await client.post(
        f"{docs_url(workspace)}/import",
        files={"file": (filename, data, DOCX_CONTENT_TYPE)},
        headers=user.headers,
    )


# --- The round trip ---


def test_every_node_and_mark_comes_back_exactly() -> None:
    every = load("every-node.json")

    imported = read(to_docx(every, "Every node"))

    assert imported.content == every
    assert imported.dropped == []
    assert imported.title == "Every node"


def test_page_breaks_are_real_word_page_breaks_and_come_back() -> None:
    every = load("every-node.json")
    data = to_docx(every, "Breaks")

    body = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode()
    assert body.count('<w:br w:type="page"/>') == kinds(every).count("pageBreak") == 1
    assert kinds(read(data).content) == kinds(every)


def test_a_version_2_edit_survives_word_and_back() -> None:
    """The fixture the realtime tests replay: an old document with a page
    break and a table added to it."""
    edited = load("v2-edit.json")["doc"]

    content = read(to_docx(edited, "Edited")).content

    assert kinds(content)[-2:] == ["pageBreak", "table"]
    assert content["content"][-1] == edited["content"][-1]


def test_lists_restart_and_keep_their_start_and_style() -> None:
    lists = {
        "type": "doc",
        "content": [
            {
                "type": "orderedList",
                "attrs": {"start": 5, "type": "a"},
                "content": [
                    {
                        "type": "listItem",
                        "content": [
                            {
                                "type": "paragraph",
                                "content": [{"type": "text", "text": "e"}],
                            }
                        ],
                    },
                    {
                        "type": "listItem",
                        "content": [
                            {
                                "type": "paragraph",
                                "content": [{"type": "text", "text": "f"}],
                            }
                        ],
                    },
                ],
            },
            {"type": "paragraph", "content": [{"type": "text", "text": "between"}]},
            {
                "type": "orderedList",
                "attrs": {"start": 1, "type": None},
                "content": [
                    {
                        "type": "listItem",
                        "content": [
                            {
                                "type": "paragraph",
                                "content": [{"type": "text", "text": "one"}],
                            }
                        ],
                    },
                ],
            },
        ],
    }

    content = read(to_docx(lists, "Lists")).content

    first, _, second = content["content"]
    assert first["attrs"] == {"start": 5, "type": "a"}
    assert second["attrs"] == {"start": 1, "type": None}


def test_the_writer_and_reader_agree_with_the_client_schema() -> None:
    """Imported nodes carry exactly the attributes the editor defines, so
    the editor never rewrites them -- and sends that as an edit -- on open."""
    nodes = manifest()["nodes"]
    marks = manifest()["marks"]
    for kind, names in attribute_names().items():
        expected = nodes[kind]["attrs"] if kind in nodes else marks[kind]
        assert names == set(expected), kind


# --- What Word writes ---


def test_words_own_list_styles_headings_and_quotes_are_understood() -> None:
    document = Document()
    document.add_heading("Title", 0)
    document.add_heading("Chapter", 2)
    document.add_paragraph("One", style="List Bullet")
    document.add_paragraph("Nested", style="List Bullet 2")
    document.add_paragraph("First", style="List Number")
    document.add_paragraph("Said someone", style="Quote")

    content = read(docx_bytes(document)).content

    assert kinds(content) == [
        "heading",
        "heading",
        "bulletList",
        "orderedList",
        "blockquote",
    ]
    assert [block["attrs"]["level"] for block in content["content"][:2]] == [1, 2]
    bullets = content["content"][2]["content"]
    assert bullets[0]["content"][1]["type"] == "bulletList"


def test_merged_cells_in_both_directions() -> None:
    document = Document()
    table = document.add_table(rows=3, cols=3)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "wide"
    table.cell(1, 2).merge(table.cell(2, 2)).text = "tall"

    rows = read(docx_bytes(document)).content["content"][0]["content"]

    spans = [
        [(c["attrs"]["colspan"], c["attrs"]["rowspan"]) for c in r["content"]]
        for r in rows
    ]
    assert spans == [[(2, 1), (1, 1)], [(1, 1), (1, 1), (1, 2)], [(1, 1), (1, 1)]]


def test_tracked_changes_keep_insertions_and_drop_deletions() -> None:
    document = body_xml(
        '<w:r><w:t xml:space="preserve">kept </w:t></w:r>'
        '<w:ins w:id="1" w:author="a"><w:r><w:t>added</w:t></w:r></w:ins>'
        '<w:del w:id="2" w:author="a"><w:r><w:delText>removed</w:delText></w:r></w:del>'
    )

    imported = read(docx_bytes(document))

    assert imported.content["content"][0]["content"] == [
        {"type": "text", "text": "kept added"}
    ]
    assert imported.dropped == [Dropped.TRACKED_CHANGE]


def test_text_inside_a_content_control_is_not_lost() -> None:
    document = body_xml(
        "<w:sdt><w:sdtContent><w:r><w:t>controlled</w:t></w:r></w:sdtContent></w:sdt>"
    )

    paragraph = read(docx_bytes(document)).content["content"][0]

    assert paragraph["content"] == [{"type": "text", "text": "controlled"}]


def test_named_highlights_and_field_codes() -> None:
    document = body_xml(
        '<w:r><w:rPr><w:highlight w:val="yellow"/></w:rPr><w:t>marked</w:t></w:r>'
        '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        "<w:r><w:instrText> PAGE </w:instrText></w:r>"
        '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
        "<w:r><w:t>7</w:t></w:r>"
        '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
    )

    inline = read(docx_bytes(document)).content["content"][0]["content"]

    assert inline[0]["marks"] == [{"type": "highlight", "attrs": {"color": "#ffff00"}}]
    # The field's instruction is not text; its last result is.
    assert inline[1] == {"type": "text", "text": "7"}


def test_a_break_inside_a_paragraph_splits_it() -> None:
    document = body_xml(
        '<w:r><w:t>before</w:t><w:br w:type="page"/><w:t>after</w:t></w:r>'
    )

    content = read(docx_bytes(document)).content

    assert kinds(content) == ["paragraph", "pageBreak", "paragraph"]


def test_what_cannot_come_is_named() -> None:
    document = Document()
    hidden = document.add_paragraph("text ").add_run("hidden")
    hidden._r.get_or_add_rPr().append(parse_xml(f"<w:vanish {nsdecls('w')}/>"))
    run = document.add_paragraph().add_run("commented")
    document.add_comment(run, text="a note")
    document.sections[0].header.paragraphs[0].text = "Running head"
    data = docx_bytes(document)
    body = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml")
    picture = b"<w:p><w:r><w:drawing/></w:r></w:p><w:p><w:r><w:t>end</w:t></w:r></w:p>"
    data = with_part(
        data, "word/document.xml", body.replace(b"<w:sectPr", picture + b"<w:sectPr")
    )

    imported = read(data)

    assert set(imported.dropped) == {
        Dropped.HIDDEN_TEXT,
        Dropped.COMMENT,
        Dropped.HEADER_FOOTER,
        Dropped.IMAGE,
    }
    # The hidden word and the picture's empty paragraph are gone.
    texts = json.dumps(imported.content)
    assert "hidden" not in texts
    assert kinds(imported.content) == ["paragraph", "paragraph", "paragraph"]


def test_every_dropped_code_has_text_in_both_languages() -> None:
    for folder in ("en", "rus"):
        path = FIXTURES.parent.parent.parent / "web" / "i18n" / folder / "document.json"
        features = json.loads(path.read_text(encoding="utf-8"))["import"]["features"]
        assert {code.value for code in Dropped} <= set(features), folder


# --- Refusals ---


@pytest.mark.parametrize(
    ("data", "status", "code"),
    [
        (compound_file("WordDocument", "1Table"), 415, "import.legacy_doc"),
        (compound_file("EncryptionInfo", "EncryptedPackage"), 422, "import.encrypted"),
        ((FIXTURES / "encrypted.docx").read_bytes(), 422, "import.encrypted"),
        (b"PK\x03\x04 not really a zip at all", 422, "import.corrupt"),
        (b"%PDF-1.7\n", 415, "import.unsupported_type"),
        (b"Plain text, named .docx", 415, "import.unsupported_type"),
        (compound_file("Workbook"), 415, "import.unsupported_type"),
        (b"", 422, "import.empty"),
    ],
    ids=[
        "doc",
        "encrypted-container",
        "encrypted-agile",
        "truncated",
        "pdf",
        "text",
        "excel-97",
        "empty",
    ],
)
async def test_a_file_that_cannot_be_imported_is_refused_with_its_own_code(
    client: AsyncClient,
    editor: TestUser,
    shared_workspace: dict[str, Any],
    data: bytes,
    status: int,
    code: str,
) -> None:
    response = await upload(client, shared_workspace, editor, data)

    assert (response.status_code, response.json()) == (status, {"code": code})
    listed = await client.get(docs_url(shared_workspace), headers=editor.headers)
    assert listed.json() == []


async def test_a_doc_is_refused_whatever_it_is_called(
    client: AsyncClient, editor: TestUser, shared_workspace: dict[str, Any]
) -> None:
    for name in ("Old report.doc", "Renamed.docx"):
        response = await upload(
            client, shared_workspace, editor, compound_file("WordDocument"), name
        )
        assert (response.status_code, response.json()["code"]) == (
            415,
            "import.legacy_doc",
        )


def test_a_damaged_document_part_is_corrupt_not_unsupported() -> None:
    data = with_part(
        docx_bytes(Document()), "word/document.xml", b"<w:document><unclosed>"
    )

    with pytest.raises(CorruptDocumentError):
        read(data)


async def test_a_file_over_the_limit_is_refused_before_it_is_read(
    client: AsyncClient,
    editor: TestUser,
    shared_workspace: dict[str, Any],
    rate_limits: Any,
) -> None:
    rate_limits(docx_import_max_bytes=1024)

    response = await upload(client, shared_workspace, editor, docx_bytes(Document()))

    assert response.status_code == 413
    assert response.json() == {"code": "import.too_large", "params": {"limit": 1024}}


# --- Endpoints and permissions ---


async def test_import_creates_a_document_named_after_the_file(
    client: AsyncClient, editor: TestUser, shared_workspace: dict[str, Any]
) -> None:
    every = load("every-node.json")

    response = await upload(
        client,
        shared_workspace,
        editor,
        to_docx(every, "Ignored"),
        "Отчёт за квартал.docx",
    )

    assert response.status_code == 201, response.text
    body = response.json()
    assert body["dropped"] == []
    assert body["document"]["title"] == "Отчёт за квартал"
    assert body["document"]["created_by"] == str(editor.id)
    stored = await client.get(
        docs_url(shared_workspace, body["document"]["id"]), headers=editor.headers
    )
    assert stored.json()["content"] == every
    assert stored.json()["version"] == 1


async def test_export_then_import_through_the_api(
    client: AsyncClient,
    owner: TestUser,
    viewer: TestUser,
    shared_workspace: dict[str, Any],
    document: dict[str, Any],
) -> None:
    every = load("every-node.json")
    saved = await client.patch(
        docs_url(shared_workspace, document["id"]),
        json={"version": 1, "content": every},
        headers=owner.headers,
    )
    assert saved.status_code == 200

    exported = await client.get(
        f"{docs_url(shared_workspace, document['id'])}/export", headers=viewer.headers
    )

    assert exported.status_code == 200
    assert exported.headers["content-type"] == DOCX_CONTENT_TYPE
    assert 'filename="Literature review.docx"' in exported.headers["content-disposition"]
    assert exported.headers["cache-control"] == "no-store"
    reimported = await upload(client, shared_workspace, owner, exported.content)
    assert reimported.json()["document"]["content"] == every


async def test_viewers_may_export_but_not_import(
    client: AsyncClient,
    viewer: TestUser,
    shared_workspace: dict[str, Any],
    document: dict[str, Any],
) -> None:
    exported = await client.get(
        f"{docs_url(shared_workspace, document['id'])}/export", headers=viewer.headers
    )
    imported = await upload(client, shared_workspace, viewer, docx_bytes(Document()))

    assert exported.status_code == 200
    assert (imported.status_code, imported.json()["code"]) == (
        403,
        "workspace.role_too_low",
    )


async def test_neither_endpoint_reaches_into_another_workspace(
    client: AsyncClient,
    owner: TestUser,
    outsider: TestUser,
    shared_workspace: dict[str, Any],
    document: dict[str, Any],
) -> None:
    exported = await client.get(
        f"{docs_url(shared_workspace, document['id'])}/export", headers=outsider.headers
    )
    imported = await upload(client, shared_workspace, outsider, docx_bytes(Document()))
    own = await client.post(
        "/workspaces", json={"name": "Elsewhere"}, headers=outsider.headers
    )
    through_own = await client.get(
        f"{docs_url(own.json(), document['id'])}/export", headers=outsider.headers
    )

    assert (exported.status_code, exported.json()["code"]) == (404, "workspace.not_found")
    assert (imported.status_code, imported.json()["code"]) == (404, "workspace.not_found")
    assert (through_own.status_code, through_own.json()["code"]) == (
        404,
        "document.not_found",
    )


async def test_both_endpoints_need_a_session(
    client: AsyncClient, shared_workspace: dict[str, Any], document: dict[str, Any]
) -> None:
    exported = await client.get(f"{docs_url(shared_workspace, document['id'])}/export")
    imported = await client.post(
        f"{docs_url(shared_workspace)}/import",
        files={"file": ("a.docx", docx_bytes(Document()), DOCX_CONTENT_TYPE)},
    )

    assert exported.status_code == imported.status_code == 401


async def test_the_body_limit_admits_an_import_larger_than_ordinary_requests(
    client: AsyncClient, editor: TestUser, shared_workspace: dict[str, Any]
) -> None:
    """Pictures make a real document large; this one carries a stored,
    incompressible part past the ordinary request limit."""
    ordinary = get_settings().max_request_body_bytes
    data = with_extra_part(docx_bytes(Document()), "word/media/scan.bin", ordinary * 2)
    assert ordinary < len(data) < get_settings().docx_import_max_bytes

    response = await upload(client, shared_workspace, editor, data)

    assert response.status_code == 201, response.text


# --- Names ---


@pytest.mark.parametrize(
    ("title", "expected"),
    [
        ("Literature review", "Literature review.docx"),
        ("Q3: plan/draft?", "Q3 plan draft.docx"),
        ("  trailing dots...  ", "trailing dots.docx"),
        ("CON", "_CON.docx"),
        ("‮evil\u0000", "evil.docx"),
        ("///", "fallback.docx"),
    ],
)
def test_the_export_file_name_is_safe_everywhere(title: str, expected: str) -> None:
    assert export_filename(title, "fallback") == expected


def test_a_non_ascii_name_is_sent_both_ways() -> None:
    header = content_disposition("Отчёт.docx", "doc-id")

    assert 'filename="doc-id.docx"' in header
    assert unquote(header.split("filename*=UTF-8''")[1]) == "Отчёт.docx"


def test_the_title_comes_from_the_file_then_the_properties_then_the_text() -> None:
    every = load("every-node.json")
    imported = read(to_docx(every, "From the properties"))

    assert choose_title("C:\\Users\\me\\Plan.DOCX", imported) == "Plan"
    assert choose_title(None, imported) == "From the properties"
    untitled = read(to_docx(every, ""))
    assert choose_title("", untitled) == "Every node and mark"


def test_numbering_that_continues_past_a_paragraph_keeps_counting() -> None:
    document = Document()
    for text in ("one", "two"):
        document.add_paragraph(text, style="List Number")
    document.add_paragraph("an aside")
    document.add_paragraph("three", style="List Number")

    first, _, second = read(docx_bytes(document)).content["content"]

    assert (first["type"], first["attrs"]["start"]) == ("orderedList", 1)
    assert (second["type"], second["attrs"]["start"]) == ("orderedList", 3)


def test_a_word_checkbox_makes_a_checklist_item() -> None:
    """As Word writes one: the box in a content control, in its own run and
    font, and the item's text after it."""
    w14 = "http://schemas.microsoft.com/office/word/2010/wordml"
    document = Document()
    for state, glyph, text in (("1", "☒", " Done"), ("0", "☐", " Open")):
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph._p.append(
            parse_xml(
                f'<w:sdt {nsdecls("w")} xmlns:w14="{w14}"><w:sdtPr><w14:checkbox>'
                f'<w14:checked w14:val="{state}"/></w14:checkbox></w:sdtPr>'
                '<w:sdtContent><w:r><w:rPr><w:rFonts w:ascii="MS Gothic" '
                f'w:hAnsi="MS Gothic"/></w:rPr><w:t>{glyph}</w:t></w:r></w:sdtContent>'
                "</w:sdt>"
            )
        )
        paragraph.add_run(text)

    items = read(docx_bytes(document)).content["content"][0]

    assert items["type"] == "taskList"
    assert [item["attrs"]["checked"] for item in items["content"]] == [True, False]
    assert items["content"][0]["content"][0]["content"] == [
        {"type": "text", "text": "Done"}
    ]


def test_section_breaks_and_page_break_before_become_page_breaks() -> None:
    document = body_xml(
        "<w:pPr><w:sectPr/></w:pPr><w:r><w:t>end of section one</w:t></w:r>",
        "<w:pPr><w:pageBreakBefore/></w:pPr><w:r><w:t>on a new page</w:t></w:r>",
        '<w:pPr><w:sectPr><w:type w:val="continuous"/></w:sectPr></w:pPr>'
        "<w:r><w:t>continuous</w:t></w:r>",
    )

    imported = read(docx_bytes(document))

    assert kinds(imported.content) == [
        "paragraph",
        "pageBreak",
        "pageBreak",
        "paragraph",
        "paragraph",
    ]
    assert imported.dropped == [Dropped.SECTION_BREAK]
