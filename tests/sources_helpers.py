"""Files to upload: real PDFs and DOCX documents built in memory, so no
binary fixtures live in the repository."""

from __future__ import annotations

import io
import uuid
import zipfile
from typing import Any

from docx import Document as WordDocument


def pdf_bytes(*pages: str | None) -> bytes:
    """A PDF with one page per argument, each showing that text in Helvetica,
    one line per `\\n`. None makes a page with no text at all, like a scan."""
    objects: list[bytes] = []

    def add(body: bytes) -> int:
        objects.append(body)
        return len(objects)

    catalog = add(b"")
    pages_id = add(b"")
    font = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    kids = []
    for text in pages:
        if text is None:
            stream = b"0.9 g 0 0 612 792 re f"
        else:
            lines = [
                line.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
                for line in text.split("\n")
            ]
            shown = b" T* ".join(b"(" + line.encode("latin-1") + b") Tj" for line in lines)
            stream = b"BT /F1 11 Tf 14 TL 72 720 Td " + shown + b" ET"
        content = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        kids.append(
            add(
                b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] "
                b"/Resources << /Font << /F1 %d 0 R >> >> /Contents %d 0 R >>"
                % (pages_id, font, content)
            )
        )
    objects[catalog - 1] = b"<< /Type /Catalog /Pages %d 0 R >>" % pages_id
    objects[pages_id - 1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (
        b" ".join(b"%d 0 R" % kid for kid in kids),
        len(kids),
    )

    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % number + body + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1))
    for offset in offsets:
        out.write(b"%010d 00000 n \n" % offset)
    out.write(
        b"trailer\n<< /Size %d /Root %d 0 R >>\nstartxref\n%d\n%%%%EOF\n"
        % (len(objects) + 1, catalog, xref)
    )
    return out.getvalue()


def docx_bytes(*paragraphs: str, table: list[list[str]] | None = None) -> bytes:
    document = WordDocument()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for row, values in zip(grid.rows, table, strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def zip_bytes(entries: dict[str, bytes], *, declared: dict[str, int] | None = None) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    raw = out.getvalue()
    if declared:
        # Rewrites the sizes the central directory claims, as a crafted bomb does.
        for name, size in declared.items():
            marker = raw.rfind(b"PK\x01\x02", 0, raw.rfind(name.encode()))
            raw = raw[: marker + 24] + size.to_bytes(4, "little") + raw[marker + 28 :]
    return raw


PNG = b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR" + b"\x00" * 64


class RecordingQueue:
    def __init__(self) -> None:
        self.enqueued: list[uuid.UUID] = []

    async def enqueue_source(self, source_id: uuid.UUID) -> None:
        self.enqueued.append(source_id)

    async def aclose(self) -> None:
        pass


def upload_body(name: str, data: bytes, content_type: str = "application/octet-stream") -> dict[str, Any]:
    return {"files": {"file": (name, data, content_type)}}
